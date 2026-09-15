# -*- coding: utf-8 -*-
"""Construye los JSON estaticos del sitio a partir de los archivos fuente en data/.
Se ejecuta una sola vez (manualmente) cada vez que cambian los excels/csv de origen.
"""
import csv
import json
import unicodedata
from pathlib import Path
from datetime import datetime

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"


def normalize(s):
    if s is None:
        return ""
    s = s.strip().upper()
    s = "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    )
    return s


# CORDOBA en el geojson de Georef viene como "Córdoba" -> normaliza a CORDOBA, coincide.
# CABA en las fuentes es "CIUDAD DE BUENOS AIRES"; en Georef es "Ciudad Autónoma de Buenos Aires".
PROVINCIA_ALIASES = {
    "CIUDAD DE BUENOS AIRES": "CIUDAD AUTONOMA DE BUENOS AIRES",
}


def canon_provincia(raw):
    n = normalize(raw)
    return PROVINCIA_ALIASES.get(n, n)


# ---------------------------------------------------------------------------
# Diputados: nomina completa (257), con el año de finalización de mandato.
# El sitio deriva de acá tanto los que renuevan en 2027 como los que
# continúan con mandato vigente (FinalizaMandato 2029).
# ---------------------------------------------------------------------------
def build_diputados():
    with open(DATA / "diputados.csv", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    out = []
    for row in rows:
        finaliza = row["FinalizaMandato"].strip()
        anio = int(finaliza.split("/")[-1])
        out.append(
            {
                "apellido": row["Apellido"].strip(),
                "nombre": row["Nombre"].strip(),
                "provincia": canon_provincia(row["Distrito"]),
                "bloque": row["Bloque"].strip(),
                "finalizaAnio": anio,
            }
        )
    print(f"Diputados (nómina completa): {len(out)}")
    print(f"  renuevan en 2027: {sum(1 for d in out if d['finalizaAnio'] == 2027)}")
    return out


# ---------------------------------------------------------------------------
# Senadores: nomina completa (72), con el año de cese legal (ya corregido).
# ---------------------------------------------------------------------------
CORRECCIONES_CESE_LEGAL = {
    ("SUAREZ", "RODOLFO ALEJANDRO"): datetime(2027, 12, 9),
    ("MANZUR", "JUAN LUIS"): datetime(2027, 12, 9),
}


def build_senadores():
    wb = openpyxl.load_workbook(DATA / "Senadores 2026.xlsx", data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(header)}

    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["APELLIDO"]] is None:
            continue
        apellido = str(row[idx["APELLIDO"]]).strip()
        nombre = str(row[idx["NOMBRE"]]).strip()
        cese_legal = row[idx["CESE LEGAL"]]

        key = (normalize(apellido), normalize(nombre))
        if key in CORRECCIONES_CESE_LEGAL:
            cese_legal = CORRECCIONES_CESE_LEGAL[key]

        out.append(
            {
                "apellido": apellido,
                "nombre": nombre,
                "provincia": canon_provincia(row[idx["PROVINCIA"]]),
                "bloque": str(row[idx["BLOQUE"]]).strip(),
                "ceseAnio": cese_legal.year,
            }
        )
    print(f"Senadores (nómina completa): {len(out)}")
    print(f"  renuevan en 2027: {sum(1 for s in out if s['ceseAnio'] == 2027)}")
    return out


# ---------------------------------------------------------------------------
# GeoJSON de provincias: se excluye el reclamo antártico de "Tierra del
# Fuego, Antártida e Islas del Atlántico Sur" (todo sub-polígono cuyo punto
# más al norte quede por debajo de -60° de latitud, el límite del Tratado
# Antártico) para que el mapa no se deforme al encuadrar el continente.
# ---------------------------------------------------------------------------
ANTARCTIC_LAT_CUTOFF = -60.0


def strip_antarctica(geojson):
    removed = 0
    for feature in geojson["features"]:
        geom = feature["geometry"]
        if geom["type"] != "MultiPolygon":
            continue
        kept = []
        for polygon in geom["coordinates"]:
            max_lat = max(pt[1] for ring in polygon for pt in ring)
            if max_lat <= ANTARCTIC_LAT_CUTOFF:
                removed += 1
                continue
            kept.append(polygon)
        geom["coordinates"] = kept
    print(f"GeoJSON: {removed} sub-polígonos antárticos removidos")
    return geojson


def build_geojson():
    geojson = json.loads((DATA / "provincias_raw.geojson").read_text(encoding="utf-8"))
    return strip_antarctica(geojson)


def main():
    diputados = build_diputados()
    senadores = build_senadores()

    (DATA / "diputados_todos.json").write_text(
        json.dumps(diputados, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (DATA / "senadores_todos.json").write_text(
        json.dumps(senadores, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    if (DATA / "provincias_raw.geojson").exists():
        geojson = build_geojson()
        (DATA / "provincias.geojson").write_text(
            json.dumps(geojson, ensure_ascii=False), encoding="utf-8"
        )

    from collections import Counter

    print("\nDiputados que renuevan en 2027, por provincia:")
    renuevan = [d for d in diputados if d["finalizaAnio"] == 2027]
    for prov, n in sorted(Counter(d["provincia"] for d in renuevan).items()):
        print(f"  {prov}: {n}")

    print("\nSenadores que renuevan en 2027, por provincia:")
    renuevan_s = [s for s in senadores if s["ceseAnio"] == 2027]
    for prov, n in sorted(Counter(s["provincia"] for s in renuevan_s).items()):
        print(f"  {prov}: {n}")


if __name__ == "__main__":
    main()
